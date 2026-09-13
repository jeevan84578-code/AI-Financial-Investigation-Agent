from __future__ import annotations

from datetime import datetime
from html import escape
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.db.models import Investigation


def _display(value: Any, fallback: str = "Not available") -> str:
    if value is None or value == "":
        return fallback
    return str(value)


def _date(value: datetime | None) -> str:
    return value.isoformat() if value else "Not available"


def _pdf_text(value: Any, fallback: str = "Not available") -> str:
    return escape(_display(value, fallback))


def generate_pdf(case: Investigation) -> BytesIO:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title=f"Helios FI Case {case.id}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="HeliosHeader", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, textColor=colors.HexColor("#0f766e"), alignment=TA_CENTER, spaceAfter=5))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#0f766e"), spaceBefore=14, spaceAfter=6))
    styles.add(ParagraphStyle(name="Body", parent=styles["BodyText"], fontSize=9, leading=13, spaceAfter=6))
    story: list[Any] = [
        Paragraph("HELIOS FI", styles["HeliosHeader"]),
        Paragraph("Financial Investigation Case Export", styles["Body"]),
        Paragraph(f"<b>Generated:</b> {_date(datetime.utcnow())}", styles["Body"]),
        Paragraph("Case overview", styles["Section"]),
        _table([
            ["Case ID", case.id, "Risk score", str(case.risk_score)],
            ["File name", case.filename, "Approval status", case.approval_status],
            ["Approved by", _display(case.approved_by), "Approval date", _date(case.approved_at)],
            ["Confidence score", _display(case.confidence_score), "Assigned to", _display(case.assigned_to)],
        ]),
        Paragraph("Executive summary", styles["Section"]),
        Paragraph(_pdf_text(case.executive_summary), styles["Body"]),
        Paragraph("Findings", styles["Section"]),
        _table([["Title", "Vendor", "Severity", "Detail"]] + [
            [finding.get("title", ""), finding.get("vendor", ""), finding.get("severity", ""), finding.get("detail", "")]
            for finding in (case.findings_json or [])
        ]),
        Paragraph("Timeline", styles["Section"]),
        _table([["Step", "Status", "Detail"]] + [
            [event.get("step", ""), event.get("status", ""), event.get("detail", "")]
            for event in (case.timeline_json or [])
        ]),
        Paragraph("AI report", styles["Section"]),
    ]
    report = case.ai_report or {}
    for title, key in [("Executive narrative", "executive_narrative"), ("Risk assessment", "risk_assessment"), ("Business impact", "business_impact")]:
        story.append(Paragraph(f"<b>{title}</b>", styles["Body"]))
        story.append(Paragraph(_pdf_text(report.get(key)), styles["Body"]))
    story.extend([
        Paragraph("Key evidence", styles["Body"]),
        Paragraph("<br/>".join(f"- {_pdf_text(item)}" for item in report.get("key_evidence", [])) or "Not available", styles["Body"]),
        Paragraph("Recommendations", styles["Section"]),
        Paragraph("<br/>".join(f"- {_pdf_text(item)}" for item in (case.recommendations or [])) or "No recommendations", styles["Body"]),
        Paragraph("Approval notes", styles["Section"]),
        Paragraph(_pdf_text(case.approval_notes), styles["Body"]),
    ])
    document.build(story)
    output.seek(0)
    return output


def _table(rows: list[list[str]]) -> Table:
    table = Table([[Paragraph(_pdf_text(cell), getSampleStyleSheet()["BodyText"]) for cell in row] for row in rows], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d7eeeb")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#123b3a")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#b7c9c9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def generate_excel(case: Investigation) -> BytesIO:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _style_sheet(summary)
    summary_rows = [
        ("Case ID", case.id), ("File name", case.filename), ("Risk score", case.risk_score),
        ("Approval status", case.approval_status), ("Approved by", _display(case.approved_by)),
        ("Approval date", _date(case.approved_at)), ("Executive summary", case.executive_summary),
        ("Confidence score", _display(case.confidence_score)), ("Generated timestamp", _date(datetime.utcnow())),
    ]
    for row in summary_rows:
        summary.append(row)
    _write_findings(workbook.create_sheet("Findings"), case)
    _write_timeline(workbook.create_sheet("Timeline"), case)
    recommendations = workbook.create_sheet("Recommendations")
    recommendations.append(["Recommendation"])
    for item in case.recommendations or []:
        recommendations.append([item])
    _style_sheet(recommendations)
    history = workbook.create_sheet("Approval History")
    history.append(["Status", "Approver", "Notes", "Timestamp"])
    for item in case.approval_history or []:
        history.append([item.get("status", ""), item.get("approver", ""), item.get("notes", ""), item.get("created_at", "")])
    _style_sheet(history)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 60)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _style_sheet(sheet: Any) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="147D78")


def _write_findings(sheet: Any, case: Investigation) -> None:
    sheet.append(["Title", "Type", "Vendor", "Severity", "Detail"])
    for item in case.findings_json or []:
        sheet.append([item.get("title", ""), item.get("type", ""), item.get("vendor", ""), item.get("severity", ""), item.get("detail", "")])
    _style_sheet(sheet)


def _write_timeline(sheet: Any, case: Investigation) -> None:
    sheet.append(["Step", "Status", "Detail"])
    for item in case.timeline_json or []:
        sheet.append([item.get("step", ""), item.get("status", ""), item.get("detail", "")])
    _style_sheet(sheet)
